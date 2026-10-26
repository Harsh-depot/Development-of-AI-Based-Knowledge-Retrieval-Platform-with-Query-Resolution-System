import os
import uuid
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
import pypdf
import docx
import pandas as pd
from langchain_text_splitters import RecursiveCharacterTextSplitter
from fastembed import TextEmbedding
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct, Filter, FieldCondition, MatchValue
from sqlalchemy.orm import Session

from config import settings
from database import Document, Chunk, get_db

# Suppress Hugging Face symlink warnings on Windows
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"

logger = logging.getLogger("ingestion")
logging.basicConfig(level=logging.INFO)


class IngestionEngine:
    def __init__(self):
        self.embedding_model = TextEmbedding(model_name=settings.FASTEMBED_MODEL_NAME)
        self.vector_dim = 384  # bge-small-en-v1.5 vector dimensionality
        self.collection_name = settings.QDRANT_COLLECTION_NAME

        # Initialize Qdrant Client (Cloud if URL provided, else persistent local disk)
        if settings.QDRANT_URL and settings.QDRANT_URL.strip():
            logger.info(f"Connecting to remote Qdrant Cloud at {settings.QDRANT_URL}")
            self.qdrant_client = QdrantClient(
                url=settings.QDRANT_URL,
                api_key=settings.QDRANT_API_KEY if settings.QDRANT_API_KEY else None,
            )
        else:
            logger.info(f"Using local persistent Qdrant at {settings.LOCAL_QDRANT_PATH}")
            os.makedirs(settings.LOCAL_QDRANT_PATH, exist_ok=True)
            self.qdrant_client = QdrantClient(path=settings.LOCAL_QDRANT_PATH)

        self._ensure_collection()

    def _ensure_collection(self):
        """Ensures the Qdrant vector collection exists."""
        try:
            collections = self.qdrant_client.get_collections().collections
            exists = any(c.name == self.collection_name for c in collections)
            if not exists:
                logger.info(f"Creating Qdrant collection '{self.collection_name}' with dim {self.vector_dim}")
                self.qdrant_client.create_collection(
                    collection_name=self.collection_name,
                    vectors_config=VectorParams(size=self.vector_dim, distance=Distance.COSINE),
                )
        except Exception as e:
            logger.error(f"Error connecting to Qdrant Cloud ({e}). Falling back to local persistent Qdrant storage.")
            try:
                os.makedirs(settings.LOCAL_QDRANT_PATH, exist_ok=True)
                self.qdrant_client = QdrantClient(path=settings.LOCAL_QDRANT_PATH)
                collections = self.qdrant_client.get_collections().collections
                if not any(c.name == self.collection_name for c in collections):
                    self.qdrant_client.create_collection(
                        collection_name=self.collection_name,
                        vectors_config=VectorParams(size=self.vector_dim, distance=Distance.COSINE),
                    )
            except Exception as e2:
                logger.error(f"Local Qdrant fallback error: {e2}")

    # =========================================================================
    # Multi-Format Parsers
    # =========================================================================

    @staticmethod
    def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
        """Parses PDF preserving page numbers."""
        pages_content = []
        reader = pypdf.PdfReader(file_path)
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_content.append({"text": text.strip(), "page_number": i + 1})
        return pages_content

    @staticmethod
    def parse_docx(file_path: str) -> List[Dict[str, Any]]:
        """Parses DOCX documents."""
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)

        joined = "\n\n".join(full_text)
        return [{"text": joined, "page_number": 1}] if joined else []

    @staticmethod
    def parse_txt(file_path: str) -> List[Dict[str, Any]]:
        """Parses plain TXT file with utf-8 / latin-1 fallback."""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1") as f:
                text = f.read()
        return [{"text": text.strip(), "page_number": 1}] if text.strip() else []

    @staticmethod
    def parse_csv(file_path: str) -> List[Dict[str, Any]]:
        """Parses FAQ or tabular CSV files into contextual question-answer units."""
        df = pd.read_csv(file_path)
        results = []
        col_names = [c.lower() for c in df.columns]

        # Check if question/answer structure exists
        q_col = next((c for c in df.columns if "question" in c.lower() or "query" in c.lower() or "prompt" in c.lower()), None)
        a_col = next((c for c in df.columns if "answer" in c.lower() or "response" in c.lower() or "resolution" in c.lower()), None)

        if q_col and a_col:
            for row_idx, (_, row) in enumerate(df.iterrows(), start=1):
                q = str(row[q_col]).strip()
                a = str(row[a_col]).strip()
                if q and a:
                    faq_text = f"FAQ Item #{row_idx}\nQuestion: {q}\nAnswer: {a}"
                    results.append({"text": faq_text, "page_number": 1})
        else:
            # General CSV: serialize row entries
            for row_idx, (_, row) in enumerate(df.iterrows(), start=1):
                row_str = " ; ".join(f"{col}: {val}" for col, val in row.items() if pd.notna(val))
                if row_str.strip():
                    results.append({"text": f"Row #{row_idx}: {row_str}", "page_number": (row_idx // 20) + 1})

        return results

    def extract_document_pages(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        """Determines format and extracts page-level text."""
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self.parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return self.parse_docx(file_path)
        elif ext == ".csv":
            return self.parse_csv(file_path)
        elif ext in [".txt", ".md"]:
            return self.parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Supported formats: PDF, DOCX, TXT, CSV")

    # =========================================================================
    # Chunking & Embedding Pipeline
    # =========================================================================

    def chunk_and_index_document(
        self,
        file_path: str,
        filename: str,
        domain_category: str,
        db_session: Session,
    ) -> Document:
        """Parses document, applies RecursiveCharacterTextSplitter (chunk size: 800, overlap: 120),

        computes FastEmbed vectors, indexes into Qdrant, and commits to SQL DB.
        """
        pages = self.extract_document_pages(file_path, filename)
        if not pages:
            raise ValueError(f"No readable text content extracted from {filename}")

        # Determine file type
        ext = Path(filename).suffix.lower().replace(".", "")

        # Create Document record
        doc_record = Document(
            filename=filename,
            file_type=ext,
            domain_category=domain_category or "General",
        )
        db_session.add(doc_record)
        db_session.flush()  # Populates doc_record.id

        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

        all_chunk_records = []
        chunk_texts_for_embedding = []
        chunk_metadata_list = []

        chunk_counter = 0
        for page in pages:
            raw_text = page["text"]
            page_num = page["page_number"]
            splits = text_splitter.split_text(raw_text)

            for split in splits:
                clean_text = split.strip()
                if not clean_text:
                    continue

                chunk_counter += 1
                chunk_id = str(uuid.uuid4())

                # SQLAlchemy record
                chunk_obj = Chunk(
                    id=chunk_id,
                    document_id=doc_record.id,
                    chunk_index=chunk_counter,
                    chunk_text=clean_text,
                    page_number=page_num,
                )
                all_chunk_records.append(chunk_obj)

                # Prepare for Qdrant
                chunk_texts_for_embedding.append(clean_text)
                chunk_metadata_list.append(
                    {
                        "chunk_id": chunk_id,
                        "document_id": doc_record.id,
                        "filename": filename,
                        "chunk_index": chunk_counter,
                        "page_number": page_num,
                        "chunk_text": clean_text,
                        "domain_category": domain_category,
                    }
                )

        if not all_chunk_records:
            raise ValueError(f"Could not generate text chunks for {filename}")

        # Batch compute local CPU FastEmbed vectors
        logger.info(f"Generating FastEmbed embeddings for {len(chunk_texts_for_embedding)} chunks...")
        embeddings = list(self.embedding_model.embed(chunk_texts_for_embedding))

        # Prepare Qdrant Points
        points = []
        for i, (embedding, meta) in enumerate(zip(embeddings, chunk_metadata_list)):
            points.append(
                PointStruct(
                    id=meta["chunk_id"],
                    vector=embedding.tolist(),
                    payload=meta,
                )
            )

        # Batch upload to Qdrant
        logger.info(f"Indexing {len(points)} points into Qdrant collection '{self.collection_name}'...")
        self.qdrant_client.upsert(
            collection_name=self.collection_name,
            points=points,
        )

        # Update SQL DB
        doc_record.chunk_count = len(all_chunk_records)
        db_session.add_all(all_chunk_records)
        db_session.commit()
        db_session.refresh(doc_record)

        logger.info(f"Successfully indexed document '{filename}' with {len(all_chunk_records)} chunks.")
        return doc_record

    # =========================================================================
    # Semantic Search
    # =========================================================================

    def search_candidates(
        self,
        query: str,
        top_k: int = None,
        domain_filter: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        """Embeds query using FastEmbed and searches Qdrant for top candidate chunks."""
        top_k = top_k or settings.RETRIEVAL_TOP_K

        # Generate query embedding
        query_embeddings = list(self.embedding_model.embed([query]))
        if not query_embeddings:
            return []
        query_vector = query_embeddings[0].tolist()

        # Optional Qdrant filter
        qdrant_filter = None
        if domain_filter and domain_filter.lower() != "all":
            qdrant_filter = Filter(
                must=[
                    FieldCondition(
                        key="domain_category",
                        match=MatchValue(value=domain_filter),
                    )
                ]
            )

        # Search in Qdrant
        try:
            # qdrant-client >= 1.10 supports query_points or search
            if hasattr(self.qdrant_client, "query_points"):
                search_result = self.qdrant_client.query_points(
                    collection_name=self.collection_name,
                    query=query_vector,
                    query_filter=qdrant_filter,
                    limit=top_k,
                ).points
            else:
                search_result = self.qdrant_client.search(
                    collection_name=self.collection_name,
                    query_vector=query_vector,
                    query_filter=qdrant_filter,
                    limit=top_k,
                )

            candidates = []
            for hit in search_result:
                payload = hit.payload or {}
                candidates.append(
                    {
                        "chunk_id": str(hit.id),
                        "score": float(hit.score),
                        "chunk_text": payload.get("chunk_text", ""),
                        "document_id": payload.get("document_id", ""),
                        "filename": payload.get("filename", "Unknown Document"),
                        "page_number": payload.get("page_number", 1),
                        "chunk_index": payload.get("chunk_index", 0),
                        "domain_category": payload.get("domain_category", "General"),
                    }
                )
            return candidates
        except Exception as e:
            logger.error(f"Error during Qdrant search: {e}")
            return []

    def delete_document(self, document_id: str, db_session: Session) -> bool:
        """Deletes a document and its chunks from both Qdrant and SQL DB."""
        try:
            # Delete points from Qdrant
            self.qdrant_client.delete(
                collection_name=self.collection_name,
                points_selector=Filter(
                    must=[
                        FieldCondition(
                            key="document_id",
                            match=MatchValue(value=document_id),
                        )
                    ]
                ),
            )

            # Delete from DB
            doc = db_session.query(Document).filter(Document.id == document_id).first()
            if doc:
                db_session.delete(doc)
                db_session.commit()
                return True
            return False
        except Exception as e:
            logger.error(f"Error deleting document {document_id}: {e}")
            db_session.rollback()
            return False


# Global singleton instance
ingestion_engine = IngestionEngine()
